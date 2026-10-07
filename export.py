# ============================================================
# export.py  —  Excel / Power BI Export Module
# ============================================================
from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter
import io

DARK   = "1F1F1F"
NAVY   = "1F3864"
RED    = "C00000"
AMBER  = "ED7D31"
GREEN  = "70AD47"
LGREY  = "F2F2F2"
WHITE  = "FFFFFF"

def _fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def _font(bold=False, color=DARK, size=11, italic=False):
    return Font(bold=bold, color=color, size=size, italic=italic,
                name="Calibri")

def _border():
    s = Side(style="thin", color="AAAAAA")
    return Border(left=s, right=s, top=s, bottom=s)

def _centre():
    return Alignment(horizontal="center", vertical="center", wrap_text=True)

def _left():
    return Alignment(horizontal="left", vertical="center", wrap_text=True)

def _risk_fill(level):
    return _fill(RED if level == "High"
                 else AMBER if level == "Medium"
                 else GREEN)


def generate_excel(results: dict, org_name: str = "Organisation") -> bytes:
    wb = Workbook()

    # ── Sheet 1: Summary ────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Summary"
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 36

    # Title
    ws.merge_cells("A1:E1")
    ws["A1"] = f"NIST CSF 2.0 Risk Assessment — {org_name}"
    ws["A1"].font      = _font(bold=True, color=WHITE, size=14)
    ws["A1"].fill      = _fill(NAVY)
    ws["A1"].alignment = _centre()
    ws.row_dimensions[1].height = 30

    # Overall score row
    ws.merge_cells("A2:C2")
    ws["A2"] = "Overall Risk Score"
    ws["A2"].font      = _font(bold=True)
    ws["A2"].fill      = _fill(LGREY)
    ws["A2"].alignment = _left()

    ws.merge_cells("D2:E2")
    ws["D2"] = f"{results['overall_score']} / 25  —  {results['overall_level']} Risk"
    ws["D2"].font      = _font(bold=True, color=WHITE)
    ws["D2"].fill      = _risk_fill(results["overall_level"])
    ws["D2"].alignment = _centre()

    # Header row
    headers = ["NIST CSF Function", "Risk Score (/25)",
               "Risk Level", "Maturity Level", "Top Recommendation"]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=4, column=col, value=h)
        c.font      = _font(bold=True, color=WHITE)
        c.fill      = _fill(NAVY)
        c.alignment = _centre()
        c.border    = _border()
    ws.row_dimensions[4].height = 22

    # Function rows
    row = 5
    for fn in results["sorted_functions"]:
        d = results["functions"][fn]
        shade = LGREY if row % 2 == 0 else WHITE

        vals = [fn, d["risk_score"], d["risk_level"],
                d["maturity"], d["recommendations"][0]]
        for col, v in enumerate(vals, 1):
            c = ws.cell(row=row, column=col, value=v)
            c.font      = _font()
            c.alignment = _left() if col in (1, 4, 5) else _centre()
            c.border    = _border()
            c.fill      = _fill(shade)
            if col == 3:
                c.fill = _risk_fill(d["risk_level"])
                c.font = _font(bold=True, color=WHITE)
        ws.row_dimensions[row].height = 48
        row += 1

    # ── Sheet 2: Function Detail ─────────────────────────────────────────
    ws2 = wb.create_sheet("Function Detail")
    ws2.column_dimensions["A"].width = 14
    ws2.column_dimensions["B"].width = 14
    ws2.column_dimensions["C"].width = 14
    ws2.column_dimensions["D"].width = 22
    ws2.column_dimensions["E"].width = 18
    ws2.column_dimensions["F"].width = 50

    ws2.merge_cells("A1:F1")
    ws2["A1"] = "Function-Level Detail"
    ws2["A1"].font      = _font(bold=True, color=WHITE, size=13)
    ws2["A1"].fill      = _fill(NAVY)
    ws2["A1"].alignment = _centre()
    ws2.row_dimensions[1].height = 28

    h2 = ["Function", "Risk Score", "Risk Level",
          "Maturity Level", "Avg Response", "Recommendations"]
    for col, h in enumerate(h2, 1):
        c = ws2.cell(row=2, column=col, value=h)
        c.font      = _font(bold=True, color=WHITE)
        c.fill      = _fill("444444")
        c.alignment = _centre()
        c.border    = _border()

    row = 3
    for fn in results["sorted_functions"]:
        d   = results["functions"][fn]
        rec = "\n".join(f"{i+1}. {r}" for i, r in
                        enumerate(d["recommendations"]))
        shade = LGREY if row % 2 == 0 else WHITE
        vals  = [fn, d["risk_score"], d["risk_level"],
                 d["maturity"], d["avg"], rec]
        for col, v in enumerate(vals, 1):
            c = ws2.cell(row=row, column=col, value=v)
            c.font      = _font(size=10)
            c.alignment = (_left() if col in (1, 4, 6) else _centre())
            c.border    = _border()
            c.fill      = _fill(shade)
            if col == 3:
                c.fill = _risk_fill(d["risk_level"])
                c.font = _font(bold=True, color=WHITE, size=10)
        ws2.row_dimensions[row].height = 90
        row += 1

    # ── Sheet 3: Raw Data (Power BI friendly) ────────────────────────────
    ws3 = wb.create_sheet("Raw Data")
    ws3.column_dimensions["A"].width = 16
    ws3.column_dimensions["B"].width = 16
    ws3.column_dimensions["C"].width = 16
    ws3.column_dimensions["D"].width = 16
    ws3.column_dimensions["E"].width = 16
    ws3.column_dimensions["F"].width = 16
    ws3.column_dimensions["G"].width = 22

    raw_headers = ["Function", "Q1", "Q2", "Q3", "Q4", "Q5",
                   "Avg Score", "Likelihood", "Impact",
                   "Risk Score", "Weight", "Weighted Score",
                   "Risk Level", "Maturity"]
    for col, h in enumerate(raw_headers, 1):
        ws3.column_dimensions[get_column_letter(col)].width = 14
        c = ws3.cell(row=1, column=col, value=h)
        c.font      = _font(bold=True, color=WHITE)
        c.fill      = _fill(NAVY)
        c.alignment = _centre()
        c.border    = _border()

    row = 2
    for fn, d in results["functions"].items():
        shade = LGREY if row % 2 == 0 else WHITE
        vals = ([fn] + d["scores"] +
                [d["avg"], d["likelihood"], d["impact"],
                 d["risk_score"], d["weight"], d["weighted"],
                 d["risk_level"], d["maturity"]])
        for col, v in enumerate(vals, 1):
            c = ws3.cell(row=row, column=col, value=v)
            c.font      = _font(size=10)
            c.alignment = _centre()
            c.border    = _border()
            c.fill      = _fill(shade)
        row += 1

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
